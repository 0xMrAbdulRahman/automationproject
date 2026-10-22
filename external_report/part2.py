"""
التقرير الخارجي - الجزء الثاني (Part 2)
Used by app.py
"""
import io

import numpy as np
import pandas as pd
import streamlit as st

from part1 import file_block, norm, ready, style_sheet

NUM_RE = r"^-?(0|[1-9]\d{0,14})(\.\d+)?$"


# ----------------------------------------------------------------------------
# Distribution of client counts per collector
# ----------------------------------------------------------------------------
def _gen_sizes(total, rng, weights, allow4):
    """Group sizes = how many collectors share the same count (1, 2, 3, or one 4)."""
    sizes, rem = [], total
    if allow4 and total >= 4 and rng.random() < 0.5:
        sizes.append(4)
        rem -= 4
    while rem > 0:
        s = min(int(rng.choice([1, 2, 3], p=weights)), rem)
        sizes.append(s)
        rem -= s
    return sizes


def make_counts(n, lo, hi, over_k, cap, allow4, rng):
    """
    One target count per collector.
    - every count is in [lo, hi] (except `over_k` collectors who may go up to `cap`)
    - the same count is shared by at most 4 collectors, and a group of 4
      happens at most once; most groups are of 3, 2 or 1
    """
    over_k = max(0, min(over_k, n, cap - hi))
    main_n = n - over_k
    pool = np.arange(lo, hi + 1)

    sizes = None
    for weights in [(0.20, 0.35, 0.45)] * 300 + [(0.10, 0.20, 0.70)] * 300:
        cand = _gen_sizes(main_n, rng, weights, allow4)
        if len(cand) <= len(pool):
            sizes = cand
            break
    if sizes is None:
        raise ValueError(
            f"الرينج ضيق: {len(pool)} رقم متاح مش كفاية لتوزيع {main_n} محصل "
            "بحيث مفيش رقم يتكرر أكتر من 3 مرات. وسّع الرينج.")

    values = rng.choice(pool, size=len(sizes), replace=False)
    counts = list(np.repeat(values, sizes))
    if over_k:
        counts += list(rng.choice(np.arange(hi + 1, cap + 1), size=over_k, replace=False))
    counts = np.array(counts)
    rng.shuffle(counts)
    return counts


# ----------------------------------------------------------------------------
# Core logic
# ----------------------------------------------------------------------------
def smart_numeric(df, skip=()):
    """Columns made only of plain numbers (no leading zeros) are saved as numbers."""
    out = df.copy()
    for c in out.columns:
        if c in skip:
            continue
        s = out[c]
        nn = s.dropna()
        if len(nn) and nn.astype(str).str.strip().str.match(NUM_RE).all():
            out[c] = pd.to_numeric(s.astype("string").str.strip(), errors="coerce")
    return out


def run_part2(base, bm, team, tm, coll, cm, lo, hi, over_k, cap, allow4, add_col, seed):
    rng = np.random.default_rng(seed)

    # --- 1. clean the base sheet --------------------------------------------
    b = pd.DataFrame({"Debtor": norm(base[bm["Debtor"]]),
                      "Collector": norm(base[bm["Collector"]])})
    start_rows = len(b)

    team_keys = norm(team[tm["Debtor"]])
    in_team = b["Debtor"].isin(set(team_keys.dropna()))
    removed_team = b[~in_team]
    b = b[in_team]

    collectors = sorted(norm(coll[cm["Collector"]]).dropna().unique())
    in_list = b["Collector"].isin(set(collectors))
    removed_list = b[~in_list]
    b = b[in_list]

    dups = int(b.duplicated("Debtor").sum())
    b = b.drop_duplicates("Debtor")

    # --- 2. target count per collector --------------------------------------
    order = list(rng.permutation(collectors))
    counts = make_counts(len(order), lo, hi, over_k, cap, allow4, rng)
    target = dict(zip(order, counts))

    # --- 3. random pick of debtors for each collector -----------------------
    by_coll = b.groupby("Collector")["Debtor"].apply(list).to_dict()
    picked, dist_rows = [], []
    for c in collectors:
        avail = by_coll.get(c, [])
        t = int(target[c])
        k = min(t, len(avail))
        if k:
            picked.extend((d, c) for d in rng.choice(avail, size=k, replace=False))
        dist_rows.append({
            "Collector": c, "Target": t, "Available": len(avail), "Selected": k,
            "Status": "OK" if len(avail) >= t else f"Shortfall (has {len(avail)})"})
    dist = pd.DataFrame(dist_rows)
    selected = pd.DataFrame(picked, columns=["Debtor", "Collector"])

    freq = (dist["Target"].value_counts().rename_axis("Count").reset_index(name="Times")
            .sort_values("Count").reset_index(drop=True))

    # --- 4. final: team sheet rows (all accounts) of the selected debtors ---
    owner = dict(zip(selected["Debtor"], selected["Collector"]))
    keep = team_keys.isin(owner.keys())
    final = team[keep].copy()
    if add_col:
        final["Assigned Collector"] = team_keys[keep].map(owner).values
    final = smart_numeric(final, skip=("Assigned Collector",)).reset_index(drop=True)

    stats = {
        "Base rows at start": start_rows,
        "Removed (Debtor not in team sheet)": int(len(removed_team)),
        "Removed (collector not in list)": int(len(removed_list)),
        "Duplicate Debtors dropped": dups,
        "Collectors": len(collectors),
        "Debtors selected": int(len(selected)),
        "Final rows (accounts)": int(len(final)),
        "Collectors with shortfall": int((dist["Status"] != "OK").sum()),
        "Max collectors sharing a count": int(freq["Times"].max()),
    }
    return final, dist, freq, selected, removed_team, removed_list, stats


def to_excel(final, dist, freq, selected, removed_team, removed_list) -> bytes:
    from openpyxl.styles import PatternFill
    orange = PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        sheets = {
            "Final Report": final,
            "Distribution": dist,
            "Count Frequency": freq,
            "Selected Debtors": selected,
            "Removed - Not in Team Sheet": removed_team,
            "Removed - Collector Not Listed": removed_list,
        }
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            style_sheet(xw.book[name], df)
        ws = xw.book["Distribution"]
        for i, status in enumerate(dist["Status"], start=2):
            if status != "OK":
                for c in range(1, len(dist.columns) + 1):
                    ws.cell(row=i, column=c).fill = orange
    return buf.getvalue()


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def render():
    st.title("التقرير الخارجي - الجزء الثاني")

    base = file_block("1) الشيت الأساسي (Debtor + Collector)", "p2_base", {
        "Debtor": ["debtor"], "Collector": ["collector"]})
    team = file_block("2) شيت الفريق (بعد التعديل)", "p2_team", {"Debtor": ["debtor"]})
    coll = file_block("3) شيت المحصلين (أسامي من غير تكرار)", "p2_coll", {
        "Collector": ["collector", "name", "محصل"]})

    st.subheader("إعدادات التوزيع")
    c1, c2, c3, c4 = st.columns(4)
    lo = c1.number_input("الحد الأدنى", 1, 100000, 250)
    hi = c2.number_input("الحد الأقصى", 1, 100000, 310)
    over_k = c3.number_input("كام محصل يعدّوا الرينج", 0, 1000, 0)
    cap = c4.number_input("أقصى رقم لهم", 1, 100000, 320)
    c5, c6, c7 = st.columns(3)
    allow4 = c5.checkbox("اسمح بمرة واحدة بس إن 4 محصلين ياخدوا نفس العدد", value=True)
    add_col = c6.checkbox("أضف عمود المحصل في التقرير النهائي", value=True)
    seed_txt = c7.text_input("Seed (اختياري - سيبه فاضي لتوزيع جديد كل مرة)")

    if st.button("شغّل الجزء الثاني", type="primary", key="p2_run"):
        if not all(ready(b) for b in (base, team, coll)):
            st.warning("ارفع التلات ملفات واختار كل الأعمدة الأول.")
        elif lo > hi or cap < hi:
            st.warning("راجع الحد الأدنى والأقصى.")
        else:
            seed = int(seed_txt) if seed_txt.strip().isdigit() else None
            with st.spinner("جاري المعالجة..."):
                try:
                    out = run_part2(base[0], base[1], team[0], team[1], coll[0], coll[1],
                                    int(lo), int(hi), int(over_k), int(cap),
                                    allow4, add_col, seed)
                    final, dist, freq, selected, rt, rl, stats = out
                    st.session_state["p2_result"] = {
                        "out": out, "xlsx": to_excel(final, dist, freq, selected, rt, rl)}
                except Exception as e:  # noqa: BLE001
                    st.session_state.pop("p2_result", None)
                    st.error(f"حصلت مشكلة: {e}")

    res = st.session_state.get("p2_result")
    if res:
        final, dist, freq, selected, rt, rl, stats = res["out"]
        st.divider()
        st.header("النتيجة")
        items = list(stats.items())
        for row in (items[:5], items[5:]):
            for col, (k, v) in zip(st.columns(len(row)), row):
                col.metric(k, f"{v:,}")

        st.download_button(
            "تحميل التقرير (Excel)", res["xlsx"], file_name="external_report_part2.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary")

        t1, t2, t3, t4, t5 = st.tabs(["Final Report", "Distribution (للمراجعة)",
                                      "Count Frequency", "Removed - Not in Team",
                                      "Removed - Collector Not Listed"])
        t1.dataframe(final.head(1000), use_container_width=True)
        t1.caption("بيعرض أول 1000 صف بس، الملف كامل في الـ Excel.")
        t2.caption("الصفوف اللي فيها Shortfall معناها المحصل معاه عملاء أقل من المطلوب، "
                   "اتاخد كل اللي معاه ومعلّم عليه.")
        t2.dataframe(dist, use_container_width=True)
        t3.caption("كل عدد اتكرر كام مرة بين المحصلين (بدل الـ Pivot).")
        t3.dataframe(freq, use_container_width=True)
        t4.dataframe(rt, use_container_width=True)
        t5.dataframe(rl, use_container_width=True)
