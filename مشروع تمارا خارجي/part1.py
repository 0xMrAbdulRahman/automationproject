"""
التقرير الخارجي - الجزء الأول (Part 1)
Used by app.py
"""
import io
from datetime import date, timedelta

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

YELLOW = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
ORANGE = PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")
HEADER_FILL = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
LOOKUP_COLS = ["Phone", "Main State", "Collector Name", "Debt Amount"]


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def norm(s: pd.Series) -> pd.Series:
    """Clean a key column: text, trimmed, no trailing '.0', blanks -> NaN."""
    s = s.astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    s = s.mask(s.fillna("") == "")
    return s.astype(object).where(s.notna(), np.nan)


def lookup(keys: pd.Series, src_keys: pd.Series, src_vals: pd.Series) -> pd.Series:
    """Excel-style lookup: first non-empty match, NaN when not found."""
    tmp = pd.DataFrame({"k": src_keys.values, "v": src_vals.values})
    tmp = tmp.dropna(subset=["k", "v"]).drop_duplicates("k")
    return keys.map(dict(zip(tmp["k"], tmp["v"])))


def to_number_if_possible(s: pd.Series) -> pd.Series:
    cleaned = s.astype("string").str.replace(",", "", regex=False).str.strip()
    num = pd.to_numeric(cleaned, errors="coerce")
    return num.where(num.notna(), s)


@st.cache_data(show_spinner=False)
def sheet_names(data: bytes, name: str):
    if name.lower().endswith(".csv"):
        return [None]
    return pd.ExcelFile(io.BytesIO(data)).sheet_names


@st.cache_data(show_spinner=False)
def load_table(data: bytes, name: str, sheet) -> pd.DataFrame:
    if name.lower().endswith(".csv"):
        df = pd.read_csv(io.BytesIO(data), dtype=str)
    else:
        df = pd.read_excel(io.BytesIO(data), sheet_name=sheet, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def file_block(label: str, key: str, fields: dict):
    """Upload a file, pick the sheet, then pick each needed column by hand."""
    st.subheader(label)
    up = st.file_uploader(label, type=["xlsx", "xlsm", "xls", "csv"], key=key,
                          label_visibility="collapsed")
    if up is None:
        return None
    data = up.getvalue()
    names = sheet_names(data, up.name)
    sheet = names[0]
    if len(names) > 1:
        sheet = st.selectbox("الشيت", names, key=f"{key}_sheet")
    df = load_table(data, up.name, sheet)
    cols = list(df.columns)
    st.caption(f"{len(df):,} صف · {len(cols)} عمود")
    with st.expander("عرض أول 5 صفوف"):
        st.dataframe(df.head(), use_container_width=True)
    ui = st.columns(len(fields))
    mapping = {}
    for box, field in zip(ui, fields):
        mapping[field] = box.selectbox(
            field, cols, index=None, placeholder="اختار العمود", key=f"{key}_{field}")
    return df, mapping


def ready(block) -> bool:
    """True when the file is uploaded and every column has been picked."""
    return block is not None and all(v is not None for v in block[1].values())


# ----------------------------------------------------------------------------
# Core logic
# ----------------------------------------------------------------------------
def build_report(port, pm, base, bm, day_files, phone_len):
    # --- portfolio, cleaned -------------------------------------------------
    p = pd.DataFrame({
        "Opportunity": norm(port[pm["Opportunity"]]),
        "Debtor": norm(port[pm["Debtor"]]),
        "Phone": norm(port[pm["Phone"]]),
        "Main State": port[pm["Main State"]].astype("string").str.strip()
                          .astype(object),
        "Debt Amount": to_number_if_possible(port[pm["Debt Amount"]]),
    })
    p["Main State"] = p["Main State"].where(p["Main State"].notna(), np.nan)

    # --- 2. new sheet: Opportunity then Debtor ------------------------------
    work = p[["Opportunity", "Debtor"]].copy()
    work = work.sort_values("Debtor", kind="stable", na_position="last")
    work = work.reset_index(drop=True)
    start_rows = len(work)

    # --- 3. phone lookup by Debtor ------------------------------------------
    work["Phone"] = lookup(work["Debtor"], p["Debtor"], p["Phone"])

    # --- 4a. phone length check (flag only) ---------------------------------
    length = work["Phone"].str.len()
    work["Phone Check"] = np.where(
        work["Phone"].isna(), "Missing phone",
        np.where(length == phone_len, "OK", "Length " + length.astype("Int64").astype(str)))

    # --- 4b + 5. phones shared by more than one Debtor -> remove them all ---
    valid = work.dropna(subset=["Phone", "Debtor"])
    per_phone = valid.groupby("Phone")["Debtor"].nunique()
    shared_phones = per_phone[per_phone > 1].index
    shared_mask = work["Phone"].isin(shared_phones)
    removed_shared = (work[shared_mask]
                      .assign(**{"Debtors on Phone": work.loc[shared_mask, "Phone"].map(per_phone)})
                      .sort_values(["Phone", "Debtor"]))
    work = work[~shared_mask].reset_index(drop=True)

    # --- 6. Main State by Debtor -------------------------------------------
    work["Main State"] = lookup(work["Debtor"], p["Debtor"], p["Main State"])

    # --- 7. Collector Name from the base sheet ------------------------------
    work["Collector Name"] = lookup(
        work["Debtor"], norm(base[bm["Debtor"]]), base[bm["Collector Name"]])

    # --- 8. Debt Amount by Opportunity (account number) ---------------------
    work["Debt Amount"] = lookup(work["Opportunity"], p["Opportunity"], p["Debt Amount"])

    # --- lookup issues: keep rows, only flag them ---------------------------
    def issues(row):
        miss = [c for c in LOOKUP_COLS if pd.isna(row[c])]
        return ", ".join(miss)
    work["Lookup Issues"] = work.apply(issues, axis=1)

    # --- 10. remove debtors that appear in the last two days' sheets --------
    seen = {}
    for label, (df, m) in day_files.items():
        for d in norm(df[m["Debtor"]]).dropna().unique():
            seen.setdefault(d, []).append(label)
    day_mask = work["Debtor"].isin(seen.keys())
    removed_days = work[day_mask].assign(
        **{"Found In": work.loc[day_mask, "Debtor"].map(lambda d: " + ".join(seen[d]))})
    work = work[~day_mask].reset_index(drop=True)

    cols = ["Opportunity", "Debtor", "Phone", "Phone Check", "Main State",
            "Collector Name", "Debt Amount", "Lookup Issues"]
    report = work[cols]
    flags = report[(report["Phone Check"] != "OK") | (report["Lookup Issues"] != "")]

    stats = {
        "Rows at start": start_rows,
        "Phones used by 2+ Debtors": int(len(shared_phones)),
        "Rows removed (shared phones)": int(len(removed_shared)),
        "Rows removed (last 2 days)": int(len(removed_days)),
        "Rows with wrong/missing phone": int((report["Phone Check"] != "OK").sum()),
        "Rows with lookup errors": int((report["Lookup Issues"] != "").sum()),
        "Final rows": int(len(report)),
    }
    return report, removed_shared, removed_days, flags, stats


def style_sheet(ws, df):
    for c in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        width = max([len(str(df.columns[c - 1]))] +
                    [len(str(v)) for v in df.iloc[:, c - 1].head(500)]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(width, 40)
    ws.freeze_panes = "A2"


def to_excel(report, removed_shared, removed_days, flags) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        sheets = {
            "Report": report,
            "Removed - Shared Phones": removed_shared,
            "Removed - Last 2 Days": removed_days,
            "Flags": flags,
        }
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            style_sheet(xw.book[name], df)

        # highlight problems on the Report + Flags sheets
        for name, df in (("Report", report), ("Flags", flags)):
            ws = xw.book[name]
            cols = list(df.columns)
            for i, (_, row) in enumerate(df.iterrows(), start=2):
                for c in LOOKUP_COLS:
                    if pd.isna(row[c]):
                        ws.cell(row=i, column=cols.index(c) + 1).fill = YELLOW
                if row["Phone Check"] != "OK" and pd.notna(row["Phone"]):
                    ws.cell(row=i, column=cols.index("Phone") + 1).fill = ORANGE
                    ws.cell(row=i, column=cols.index("Phone Check") + 1).fill = ORANGE
    return buf.getvalue()


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def render():
    st.title("التقرير الخارجي - الجزء الأول")

    today = date.today()
    d1, d2 = today - timedelta(days=1), today - timedelta(days=2)

    port = file_block("1) ملف المحفظة (Portfolio)", "port", {
        "Debtor": ["debtor"],
        "Opportunity": ["opportunity", "account"],
        "Phone": ["phone", "mobile", "tel"],
        "Main State": ["state", "status"],
        "Debt Amount": ["debt", "amount", "balance"],
    })
    base = file_block("2) الشيت الأساسي (Debtor + Collector)", "base", {
        "Debtor": ["debtor"],
        "Collector Name": ["collector"],
    })
    day1 = file_block(f"3) شيت يوم {d1:%d} (أمس)", "day1", {"Debtor": ["debtor"]})
    day2 = file_block(f"4) شيت يوم {d2:%d} (أول أمس)", "day2", {"Debtor": ["debtor"]})

    phone_len = st.number_input("طول رقم التليفون المطلوب", 1, 20, 10)

    if st.button("شغّل التقرير", type="primary"):
        if not all(ready(b) for b in (port, base, day1, day2)):
            st.warning("ارفع الأربع ملفات واختار كل الأعمدة الأول.")
        else:
            with st.spinner("جاري المعالجة..."):
                try:
                    report, rs, rd, flags, stats = build_report(
                        port[0], port[1], base[0], base[1],
                        {f"Day {d1:%d}": day1, f"Day {d2:%d}": day2},
                        int(phone_len))
                    st.session_state["result"] = {
                        "report": report, "rs": rs, "rd": rd, "flags": flags,
                        "stats": stats,
                        "xlsx": to_excel(report, rs, rd, flags),
                    }
                except Exception as e:  # noqa: BLE001
                    st.session_state.pop("result", None)
                    st.error(f"حصلت مشكلة: {e}")

    res = st.session_state.get("result")
    if res:
        st.divider()
        st.header("النتيجة")
        cols = st.columns(len(res["stats"]))
        for c, (k, v) in zip(cols, res["stats"].items()):
            c.metric(k, f"{v:,}")

        st.download_button(
            "تحميل التقرير (Excel)", res["xlsx"],
            file_name=f"report_{today:%Y-%m-%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary")

        t1, t2, t3, t4 = st.tabs(["Report", "Flags (للمراجعة)",
                                  "Removed - Shared Phones", "Removed - Last 2 Days"])
        t1.dataframe(res["report"], use_container_width=True)
        t2.caption("أصفر = Lookup ملقاش قيمة · برتقالي = طول التليفون غلط. الصفوف دي لسه في الداتا.")
        t2.dataframe(res["flags"], use_container_width=True)
        t3.dataframe(res["rs"], use_container_width=True)
        t4.dataframe(res["rd"], use_container_width=True)

